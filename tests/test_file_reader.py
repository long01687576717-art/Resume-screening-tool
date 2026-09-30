"""回归测试：read_bytes（读内存，网页"原简历"用）和 read_file（读文件，分析用）结果必须完全一致。

两者共用同一套 TXT 解码、Word 遍历和收尾处理；这个测试防止以后只改了其中一边。
覆盖：samples/ 里的虚构 TXT 和 Word、一份现场生成的复杂 Word（合并单元格、表格套表格、
多种样式、文本框）、TXT 的各种编码（含现有读取器读不好的 UTF-16、GB18030，要求两边"错得一样"）。
另外检查 pdf_pages（"原简历"里 PDF 转图片）的清晰度：A4 / Letter 渲染成 1460 像素宽，小页面分辨率不超过 200 DPI。

用法（在项目根目录）：python tests/test_file_reader.py      也可以用 pytest 运行
不需要浏览器、API Key；生成的文件放在系统临时目录，结束后自动删除。全部通过退出码 0，否则 1。
"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                     # noqa: E402
from docx.oxml import parse_xml               # noqa: E402
from docx.shared import RGBColor              # noqa: E402
from PIL import Image                         # noqa: E402

from parser.file_reader import PAGE_IMAGE_WIDTH, pdf_pages, read_bytes, read_file   # noqa: E402

SAMPLE_TEXT = "张三（虚构样例）\n教育：某某大学 统计学\n技能：Python、SQL\n生僻字：𠀀 喆 昇"
ENCODINGS = {
    "utf8": SAMPLE_TEXT.encode("utf-8"),
    "utf8_bom": b"\xef\xbb\xbf" + SAMPLE_TEXT.encode("utf-8"),
    "gbk": SAMPLE_TEXT.replace("𠀀", "").encode("gbk"),
    "gb18030": SAMPLE_TEXT.encode("gb18030"),          # 现有读取器会读成乱码（已知局限），两边要一致
    "utf16_bom": SAMPLE_TEXT.encode("utf-16"),         # 同上
    "utf16be": SAMPLE_TEXT.encode("utf-16-be"),        # 同上
    "bad_byte": SAMPLE_TEXT.encode("utf-8")[:20] + b"\xff\xfe" + SAMPLE_TEXT.encode("utf-8")[20:],
    "empty": b"",                                      # 两边都应报"没有识别到文字"
}


def _complex_docx(path):
    doc = Document()
    doc.add_heading("李想（虚构样例）", level=1)
    p = doc.add_paragraph()
    p.add_run("求职意向：").bold = True
    run = p.add_run("数据分析")
    run.italic = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "教育背景"
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(1, 0).text = "2022.09-2026.06"
    table.cell(1, 1).text = "某某大学"
    table.cell(1, 2).text = "统计学 本科"
    inner = table.cell(1, 2).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "GPA"
    inner.cell(0, 1).text = "3.6/4.0"
    doc.add_paragraph("负责周报自动化，节省 2 小时/周", style="List Bullet")
    box = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t>文本框：技能 Python、SQL</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
    doc.element.body.insert(len(doc.element.body) - 1, box)
    doc.save(path)


def _read(fn, *args):
    """返回 ("ok", 文字) 或 ("error", 报错信息)，这样"都报同样的错"也算一致。"""
    try:
        result = fn(*args)
        return "ok", result[0] if isinstance(result, tuple) else result
    except ValueError as e:
        return "error", str(e)


def _cases(tmp):
    files = sorted((ROOT / "samples").rglob("*.txt")) + sorted((ROOT / "samples").rglob("*.docx"))
    files = [f for f in files if "jd" not in f.parts]               # samples/jd 是岗位描述，不是简历
    generated = tmp / "复杂排版.docx"
    _complex_docx(generated)
    files.append(generated)
    for name, data in ENCODINGS.items():
        path = tmp / f"编码_{name}.txt"
        path.write_bytes(data)
        files.append(path)
    return files


def check_all():
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        for path in _cases(Path(tmp)):
            from_file = _read(read_file, path)
            from_mem = _read(read_bytes, path.name, path.read_bytes())
            detail = ""
            if from_file != from_mem:
                a, b = from_file[1].splitlines(), from_mem[1].splitlines()
                first = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), min(len(a), len(b)))
                detail = f"第 {first + 1} 行起不同：文件 {a[first:first + 1]} / 内存 {b[first:first + 1]}"
            results.append((path.name, from_file == from_mem, from_file[0], detail))
        generated_text = _read(read_bytes, "复杂排版.docx", (Path(tmp) / "复杂排版.docx").read_bytes())[1]
    must_have = ["教育背景", "3.6/4.0", "文本框：技能", "负责周报自动化"]
    results.append(("复杂 Word 的表格、嵌套表格、文本框都读到", all(k in generated_text for k in must_have),
                    "ok", f"缺少：{[k for k in must_have if k not in generated_text]}"))
    for suffix in (".pdf", ".png"):
        status, message = _read(read_bytes, f"x{suffix}", b"")
        results.append((f"read_bytes 拒绝 {suffix}", status == "error" and "只支持" in message, status, message))
    results += _check_pdf_pages()
    return results


def _pdf(width_pt, height_pt):
    """用 PIL 生成一页指定大小（单位：点，1/72 英寸）的 PDF，放在内存里。"""
    buf = io.BytesIO()
    Image.new("RGB", (int(width_pt * 2), int(height_pt * 2)), "white").save(buf, format="PDF", resolution=144)
    return buf.getvalue()


def _check_pdf_pages():
    out = []
    for label, size, expect in (("A4", (595, 842), PAGE_IMAGE_WIDTH), ("Letter", (612, 792), PAGE_IMAGE_WIDTH),
                                ("小页面 A6（封顶 200 DPI）", (298, 420), round(298 * 200 / 72))):
        width = pdf_pages(_pdf(*size))[0].size[0]
        out.append((f"PDF 转图片 {label}：宽 {expect} 像素左右", abs(width - expect) <= 2, "", f"实际宽 {width}"))
    return out


def test_read_bytes_matches_read_file():
    failed = [r for r in check_all() if not r[1]]
    assert not failed, failed


def main():
    results = check_all()
    for name, ok, status, detail in results:
        kind = {"ok": "（读出文字）", "error": "（报错）"}.get(status, "")
        print(f"  [{'通过' if ok else '失败'}] {name}{kind}" + (f"：{detail}" if not ok else ""))
    passed = all(r[1] for r in results)
    print(f"共 {len(results)} 项，" + ("全部通过" if passed else f"{sum(not r[1] for r in results)} 项失败"))
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())
